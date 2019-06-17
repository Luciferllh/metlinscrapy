# @Time    : 2019/6/17 13:08
# @Email   : 772632967@qq.com
# @File    : word.py
# @Software: PyCharm
from docx import Document



class WordMass(object):
    """
    提取word文件中的mass

    """
    def __init__(self,path):
        self.path=path
        self.mass_list=self.get_mass_list()

    def get_mass_list(self):
        document = Document(self.path)
        tables = document.tables
        table = tables[0]
        mass_list = []
        for i in range(1, len(table.rows)):
            result = table.cell(i, 0).text
            # cell(i,0)表示第(i+1)行第1列数据，以此类推
            mass_list.append(result)
        return mass_list

